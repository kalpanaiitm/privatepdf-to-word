from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from privatepdf.converter import ConversionError, build_docx, convert_pdf_bytes
from privatepdf.sample import make_sample_pdf


def test_sample_pdf_converts_to_valid_docx():
    result = convert_pdf_bytes(make_sample_pdf())
    assert result.page_count == 1
    assert result.character_count > 100
    doc = Document(BytesIO(result.docx_bytes))
    assert "Fictional Project Summary" in "\n".join(p.text for p in doc.paragraphs)
    assert len([p for p in doc.paragraphs if p.text.strip()]) >= 5
    assert any(p.text == "Fictional Project Summary" and p.style.name.startswith("Heading") for p in doc.paragraphs)


def test_non_pdf_is_rejected():
    with pytest.raises(ConversionError, match="valid PDF signature"):
        convert_pdf_bytes(b"This is not a PDF")


def test_empty_file_is_rejected():
    with pytest.raises(ConversionError, match="empty"):
        convert_pdf_bytes(b"")


def test_encrypted_pdf_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ConversionError, match="encrypted"):
        convert_pdf_bytes(output.getvalue())


def test_scan_only_pdf_is_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ConversionError, match="scanned or image-based"):
        convert_pdf_bytes(output.getvalue())


def test_docx_builder_preserves_page_text():
    output = build_docx(["1 Introduction\n\nFirst paragraph.", "SECOND PAGE\n\nMore text."])
    doc = Document(BytesIO(output))
    content = "\n".join(p.text for p in doc.paragraphs)
    assert "First paragraph." in content
    assert "More text." in content
