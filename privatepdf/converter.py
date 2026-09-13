import re
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.enum.text import WD_BREAK

from .validation import ValidationError, validate_pdf


class ConversionError(ValueError):
    """Safe, user-facing conversion failure."""


@dataclass(frozen=True)
class ConversionResult:
    docx_bytes: bytes
    page_count: int
    character_count: int
    warnings: tuple[str, ...]


def _looks_like_heading(line: str) -> bool:
    words = line.split()
    return 1 <= len(words) <= 10 and len(line) <= 90 and (
        line.isupper()
        or line.istitle()
        or re.match(r"^\d+(?:\.\d+)*\s+\S+", line) is not None
    )


def build_docx(pages: list[str]) -> bytes:
    document = Document()
    for page_number, text in enumerate(pages, start=1):
        if page_number > 1:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        for raw_line in text.splitlines():
            line = re.sub(r"[ \t]+", " ", raw_line).strip()
            if not line:
                continue
            if _looks_like_heading(line):
                document.add_heading(line, level=2)
            else:
                document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def convert_pdf_bytes(data: bytes) -> ConversionResult:
    try:
        reader = validate_pdf(data)
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except ValidationError as exc:
        raise ConversionError(str(exc)) from exc
    except Exception as exc:
        raise ConversionError("Text extraction failed for this PDF.") from exc

    character_count = sum(len(page) for page in pages)
    if character_count < max(20, len(pages) * 5):
        raise ConversionError(
            "Very little selectable text was found. This appears to be scanned or image-based; OCR is not included."
        )
    warnings = (
        "Review the output carefully: complex layouts, tables, equations, footnotes, and images may not be preserved.",
    )
    return ConversionResult(
        docx_bytes=build_docx(pages),
        page_count=len(pages),
        character_count=character_count,
        warnings=warnings,
    )
